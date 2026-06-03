"""Build Week 15 PPTX with the real SYSU official blue template.

This script intentionally does not compress ``script.md`` into slides. It uses
the reviewed Week 15 storyboard plus a fixed template mapping, then fills a copy
of the official SYSU blue PPTX template.
"""

from __future__ import annotations

import json
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
WEEK_DIR = ROOT / "course" / "weeks" / "week_15"
EVAL_PATH = ROOT / "course" / "evaluation" / "week_15_ppt_evidence_review.md"
STORYBOARD_PATH = WEEK_DIR / "ppt_storyboard.md"
AI_PPT_ROOT = Path(r"E:\Codex_Projects\AI_PPT")
SOURCE_TEMPLATE = AI_PPT_ROOT / "templates" / "source" / "sysu-official" / "中山大学幻灯片模板-蓝.pptx"
STYLE_SPEC = AI_PPT_ROOT / "templates" / "styles" / "strict-sysu-official-blue" / "style.json"
ASSET_MANIFEST = AI_PPT_ROOT / "templates" / "assets" / "strict-sysu-official-blue" / "asset-manifest.json"
OUT_DIR = ROOT / "outputs" / "ppt" / "sysu_official_blue" / "week_15"

STYLE_ID = "strict-sysu-official-blue"
COURSE_NAME = "医药数据处理与可视化"
WEEK_TITLE = "差异表达分析与功能解读"
FONT = "思源黑体 CN Medium"
FONT_HEAVY = "思源黑体 CN Heavy"
BODY_FONT = "等线"


class P:
    accent = "0B4F6C"
    secondary = "3494BA"
    mint = "58B6C0"
    green = "75BDA7"
    slate = "373545"
    pale = "CEDBE6"
    surface = "F4F8FA"
    border = "B7CBCD"
    text = "1D2733"
    muted = "667085"
    white = "FFFFFF"
    warning = "DE8F05"
    danger = "B42318"


@dataclass(frozen=True)
class SlideSpec:
    no: int
    action_title: str
    layout_name: str
    layout_idx: int
    role: str
    exhibit: str
    teacher_note: str
    student_action: str
    risk_note: str
    bullets: list[str] = field(default_factory=list)
    body: str = ""
    render: str = "bullets"


SLIDES: list[SlideSpec] = [
    SlideSpec(
        1,
        "差异表达分析要同时看显著性、效应量和证据边界",
        "标题幻灯片",
        0,
        "cover",
        "官方蓝模板封面",
        "今天的目标不是背软件菜单，而是训练从结果表到谨慎解释的流程。",
        "说出本周会同时检查哪三类证据。",
        "不把差异表达等同于机制证明。",
        ["统计显著", "效应量大小", "证据边界"],
        "Week 15 · SYSU 官方蓝模板样板",
        "cover",
    ),
    SlideSpec(
        2,
        "Week 15 接在 count matrix 之后，服务功能解释之前",
        "节标题",
        2,
        "transition",
        "课程链路图",
        "先回忆 Week 14 的 count matrix 和 metadata，再说明本周只讨论 bulk RNA-seq 结果表。",
        "在链路图上标出 Week 15 的输入和输出。",
        "不混入单细胞 marker 或 UMAP 解释。",
        ["Week 14：count matrix 与 metadata", "Week 15：DESeq2 结果表与图形解释", "Week 16：单细胞可视化解释边界"],
        "本页用于把本周放回课程主线。",
        "chain",
    ),
    SlideSpec(
        3,
        "药物处理组和对照组的表达变化先是数据问题",
        "两栏内容",
        3,
        "scenario",
        "输入、规则、输出三问",
        "把医学问题先拆成数据字段，而不是直接问 AI 哪个基因重要。",
        "写出 treatment、control、candidate gene 三个要素。",
        "不宣称药物疗效或疾病机制。",
        ["处理组和对照组是什么", "表达量来自哪些样本", "候选基因如何进入后续核验"],
        "药物处理组 vs 对照组只是课堂场景，不能直接推出疗效。",
        "three_questions",
    ),
    SlideSpec(
        4,
        "count matrix 与 metadata 对齐决定后续分析是否可信",
        "两栏内容",
        3,
        "input-check",
        "小型 count matrix 与 metadata 对齐示意",
        "如果样本顺序或分组错了，后面所有统计和图都没有意义。",
        "检查两个表的 sample_id 是否一一对应。",
        "示例数据为教学模拟，不代表真实项目。",
        ["样本 ID 必须一致", "分组信息必须明确", "表格顺序错误会污染全流程"],
        "教学模拟，不代表真实医学结论。",
        "matrix_metadata",
    ),
    SlideSpec(
        5,
        "DESeq2 流程从标准化走向模型检验",
        "内容与标题",
        7,
        "method",
        "DESeq2 概念流程图",
        "只讲流程位置，不推导负二项模型。",
        "指出 size factor 和 dispersion 位于哪一步。",
        "模型细节需进一步教材或文献核验。",
        ["counts", "normalization", "dispersion", "model test", "results table"],
        "本页只讲本科课堂所需流程定位。",
        "process",
    ),
    SlideSpec(
        6,
        "结果表字段分别回答不同问题",
        "标题和内容",
        1,
        "field-table",
        "DESeq2 字段解释表",
        "每一列回答的问题不同，不能只盯一个数字。",
        "给每个字段写一句“它回答什么”。",
        "不把单个字段写成完整结论。",
        ["gene", "baseMean", "log2FoldChange", "pvalue", "padj"],
        "字段解释来自 Week 15 讲义与课堂小表。",
        "field_table",
    ),
    SlideSpec(
        7,
        "`padj` 和 `log2FoldChange` 共同筛选候选基因",
        "标题和内容",
        1,
        "screening",
        "4 行 DESeq2 结果小表和阈值标记",
        "先按规则筛选，再讨论为什么只是候选。",
        "手工标记上调、下调和不确定基因。",
        "阈值是教学规则，不是通用医学标准。",
        ["padj < 0.05", "abs(log2FoldChange) > 1", "先筛选，再核验"],
        "教学模拟，不代表真实医学结论。",
        "screen_table",
    ),
    SlideSpec(
        8,
        "多重检验解释为什么不能只看原始 P 值",
        "比较",
        4,
        "statistics",
        "pvalue 与 padj 判断对照",
        "同时检验大量基因会放大假阳性风险。",
        "用一句话解释为什么关注 padj。",
        "不展开未讲授的校正算法细节。",
        ["只看 pvalue：更容易留下假阳性", "加入 padj：更适合大量基因同时检验", "仍需结合效应量和表达水平"],
        "本课堂关注校正后显著性以降低大量检验中的假阳性风险。",
        "compare_padj",
    ),
    SlideSpec(
        9,
        "火山图把效应量和显著性放在同一张图里",
        "图片与标题",
        8,
        "visualization",
        "模拟火山图",
        "火山图是二维判读工具，不是机制图。",
        "写一句火山图图注。",
        "模拟点不对应真实基因功能。",
        ["横轴：log2FoldChange", "纵轴：-log10(padj)", "颜色：候选上调/下调/未通过阈值"],
        "教学模拟，不代表真实医学结论。",
        "volcano",
    ),
    SlideSpec(
        10,
        "学生先手工标记上调、下调和不确定基因",
        "标题和内容",
        1,
        "activity",
        "阈值判读练习卡",
        "工具执行前先做人工判读，暴露规则理解错误。",
        "小组给 GENE_A 到 GENE_D 贴标签。",
        "基因名为教学代号，不引用真实功能。",
        ["GENE_A：候选上调", "GENE_B：不确定", "GENE_C：不确定", "GENE_D：候选下调"],
        "教学模拟，不代表真实医学结论。",
        "label_cards",
    ),
    SlideSpec(
        11,
        "热图展示表达模式，不替代统计检验",
        "图片与标题",
        8,
        "visualization",
        "模拟热图",
        "热图看的是模式和聚类，不能替代 DESeq2 统计检验。",
        "说出样本是否大致按分组聚集。",
        "不把聚类模式写成因果或疗效。",
        ["看表达模式", "看样本是否大致聚集", "不替代统计检验"],
        "教学模拟，不代表真实医学结论。",
        "heatmap",
    ),
    SlideSpec(
        12,
        "功能解读只能从候选基因走向候选证据",
        "内容与标题",
        7,
        "interpretation",
        "候选证据流程图",
        "结果表之后进入的是证据收集，不是直接下生物学结论。",
        "列出功能解读需要核验的三类信息。",
        "未核验前只能说候选解释。",
        ["candidate genes", "database query", "literature check", "conservative report"],
        "功能解释必须标注需数据库或文献核验。",
        "evidence_flow",
    ),
    SlideSpec(
        13,
        "GO、Reactome、STRING 不能直接证明机制",
        "比较",
        4,
        "database-boundary",
        "数据库边界三列对照",
        "数据库结果是线索，不是机制证明。",
        "给每列写一个“能说”和“不能说”。",
        "数据库解释必须标注需数据库或文献核验。",
        ["GO：候选功能术语", "Reactome：候选通路线索", "STRING：候选相互作用网络"],
        "所有数据库解释均为候选证据。",
        "database_compare",
    ),
    SlideSpec(
        14,
        "AI 可以整理规则，不能编造基因功能",
        "两栏内容",
        3,
        "ai-boundary",
        "Prompt 与 AI 输出审计清单",
        "AI 只做候选规则、图注、核验清单，不替代来源核验。",
        "修改 Prompt，加入不编造和需核验约束。",
        "AI 输出不作为权威结论。",
        ["不要编造基因功能", "所有功能解释标注需核验", "输出为候选模板，不是结论"],
        "AI 只能生成候选解释、候选代码、候选图注或流程建议。",
        "ai_prompt",
    ),
    SlideSpec(
        15,
        "常见错误集中在阈值、图形和过度解释",
        "标题和内容",
        1,
        "misconception",
        "错误清单和纠偏动作表",
        "把常见错误现场纠偏，避免学生形成错误汇报习惯。",
        "选一条错误，写出改正后的表达。",
        "不把错误示例写成推荐说法。",
        ["只看 pvalue", "把 log2FC 大写成医学意义大", "把火山图当机制图", "让 AI 编造功能"],
        "错误示例只用于纠偏。",
        "mistake_table",
    ),
    SlideSpec(
        16,
        "小组任务要求提交筛选、图注和核验记录",
        "标题和内容",
        1,
        "class-task",
        "提交物清单",
        "收过程记录，不只收最终答案。",
        "提交筛选表、火山图图注、待核验列表。",
        "未提交核验记录视为不完整。",
        ["筛选表", "火山图图注", "功能候选解释", "待核验列表"],
        "课堂提交物来自 Week 15 课堂任务。",
        "submission",
    ),
    SlideSpec(
        17,
        "Claim-Evidence Gate 决定结论能否写入报告",
        "内容与标题",
        7,
        "evidence-gate",
        "证据门流程",
        "每一句结论都要能回答“证据从哪里来”。",
        "把一条结论改写为候选解释并标注证据状态。",
        "机制、通路、药效主张必须保留核验边界。",
        ["claim", "evidence status", "allowed wording", "action"],
        "Claim-Evidence Gate 来自 Week 15 讲义和 evidence review。",
        "claim_gate",
    ),
    SlideSpec(
        18,
        "出口卡收束：哪一句结论必须标注“需核验”",
        "节标题",
        2,
        "closing",
        "出口卡和下周衔接",
        "用出口卡确认学生掌握证据边界，下周转向单细胞可视化解释。",
        "写一条需核验结论并说明原因。",
        "不以“谢谢”页替代学习收束。",
        ["出口卡：哪一句结论必须标注需核验？", "理由：证据来自哪里？还缺什么？", "下周：单细胞转录组可视化"],
        "课程收束来自 Week 15 讲义。",
        "exit_card",
    ),
]


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def ensure_inputs() -> None:
    missing = [path for path in [STORYBOARD_PATH, EVAL_PATH, SOURCE_TEMPLATE, STYLE_SPEC, ASSET_MANIFEST] if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing required input(s): " + ", ".join(str(path) for path in missing))
    review = EVAL_PATH.read_text(encoding="utf-8")
    if not re.search(r"^status:\s*pass\s*$", review, flags=re.M) or not re.search(r"^pass\s*$", review, flags=re.M):
        raise RuntimeError("Evidence review must be status/pass before PPTX generation.")
    storyboard_titles = parse_storyboard_titles()
    expected = {slide.no: slide.action_title for slide in SLIDES}
    if storyboard_titles != expected:
        raise RuntimeError("Storyboard action titles do not match generator mapping.")


def parse_storyboard_titles() -> dict[int, str]:
    text = STORYBOARD_PATH.read_text(encoding="utf-8")
    found: dict[int, str] = {}
    for line in text.splitlines():
        match = re.match(r"^\|\s*(\d+)\s*\|\s*(.+?)\s*\|", line)
        if not match:
            continue
        found[int(match.group(1))] = match.group(2).strip()
    return found


def remove_all_slides(prs: Presentation) -> None:
    slide_ids = list(prs.slides._sldIdLst)  # pylint: disable=protected-access
    for slide_id in slide_ids:
        prs.part.drop_rel(slide_id.rId)
        prs.slides._sldIdLst.remove(slide_id)  # pylint: disable=protected-access


def set_text_frame(
    shape,
    text: str,
    *,
    size: int = 20,
    color: str = P.text,
    bold: bool = False,
    font: str = FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    para = tf.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)


def add_textbox(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: int = 18,
    color: str = P.text,
    bold: bool = False,
    font: str = BODY_FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text_frame(box, text, size=size, color=color, bold=bold, font=font, align=align, valign=valign)
    return box


def add_bullets(slide, items: list[str], x: float, y: float, w: float, h: float, *, size: int = 18, color: str = P.text):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    for i, item in enumerate(items):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = item
        para.level = 0
        para.font.name = BODY_FONT
        para.font.size = Pt(size)
        para.font.color.rgb = rgb(color)
        para.space_after = Pt(8)
    return box


def add_rect(slide, x: float, y: float, w: float, h: float, fill: str, line: str | None = None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_line(slide, x1: float, y1: float, x2: float, y2: float, color: str = P.accent, width: int = 2):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    return line


def add_note(slide, text: str) -> None:
    add_textbox(slide, text, 0.92, 6.58, 11.3, 0.28, size=10, color=P.muted)


def set_common_placeholders(slide, spec: SlideSpec) -> None:
    display_title = spec.action_title.replace("`", "")
    for placeholder in slide.placeholders:
        idx = placeholder.placeholder_format.idx
        ptype = placeholder.placeholder_format.type
        if ptype in {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE} or idx == 0:
            title_size = 23 if len(display_title) > 24 else 26
            set_text_frame(placeholder, display_title, size=title_size, bold=True, font=FONT_HEAVY)
        elif ptype == PP_PLACEHOLDER.DATE or idx == 10:
            set_text_frame(placeholder, "2026", size=8, color=P.muted)
        elif ptype == PP_PLACEHOLDER.FOOTER or idx == 11:
            set_text_frame(placeholder, f"{COURSE_NAME} · Week 15", size=8, color=P.muted, align=PP_ALIGN.CENTER)
        elif ptype == PP_PLACEHOLDER.SLIDE_NUMBER or idx == 12:
            set_text_frame(placeholder, f"{spec.no:02d}", size=8, color=P.muted, align=PP_ALIGN.RIGHT)
        else:
            set_text_frame(placeholder, "", size=8)


def add_source_risk(slide, spec: SlideSpec) -> None:
    add_textbox(slide, f"来源：Week 15 script/outline/materials；风险边界：{spec.risk_note}", 0.92, 6.92, 10.8, 0.18, size=8, color=P.muted)


def add_label(slide, text: str, x: float, y: float, w: float, color: str = P.accent) -> None:
    add_rect(slide, x, y, w, 0.35, color)
    add_textbox(slide, text, x + 0.08, y + 0.07, w - 0.16, 0.18, size=10, color=P.white, bold=True, align=PP_ALIGN.CENTER)


def add_card(slide, title: str, body: str, x: float, y: float, w: float, h: float, color: str = P.secondary) -> None:
    add_rect(slide, x, y, w, h, P.surface, line=P.border)
    add_rect(slide, x, y, 0.09, h, color)
    add_textbox(slide, title, x + 0.24, y + 0.18, w - 0.38, 0.28, size=15, bold=True, color=P.text)
    add_textbox(slide, body, x + 0.24, y + 0.62, w - 0.38, h - 0.75, size=13, color=P.muted)


def add_simple_table(slide, rows: list[list[str]], x: float, y: float, w: float, h: float, header_fill: str = P.accent) -> None:
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for c in range(len(rows[0])):
        table.columns[c].width = Inches(w / len(rows[0]))
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = text
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(header_fill if r == 0 else P.white)
            for para in cell.text_frame.paragraphs:
                para.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
                for run in para.runs:
                    run.font.name = BODY_FONT
                    run.font.size = Pt(9 if r == 0 else 10)
                    run.font.bold = r == 0
                    run.font.color.rgb = rgb(P.white if r == 0 else P.text)


def render_cover(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    for placeholder in slide.placeholders:
        if placeholder.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE or placeholder.placeholder_format.idx == 1:
            set_text_frame(placeholder, f"{WEEK_TITLE}\n{spec.body}", size=22, color=P.accent, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, "本周主线", 0.92, 5.45, 1.2, 0.25, size=12, color=P.secondary, bold=True)
    add_bullets(slide, spec.bullets, 2.1, 5.38, 7.2, 0.8, size=15, color=P.text)


def render_chain(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    xs = [1.0, 4.7, 8.4]
    colors = [P.secondary, P.accent, P.green]
    labels = ["Week 14", "Week 15", "Week 16"]
    for i, (x, label, color, text) in enumerate(zip(xs, labels, colors, spec.bullets)):
        add_rect(slide, x, 2.6, 2.7, 1.3, P.surface, line=P.border)
        add_label(slide, label, x, 2.6, 2.7, color)
        add_textbox(slide, text.split("：", 1)[-1], x + 0.2, 3.12, 2.25, 0.38, size=15, color=P.text, bold=True, align=PP_ALIGN.CENTER)
        if i < 2:
            add_line(slide, x + 2.85, 3.25, xs[i + 1] - 0.15, 3.25, color=P.border, width=3)
    add_note(slide, spec.teacher_note)


def render_three_questions(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    left = ["处理组", "对照组", "候选基因"]
    right = ["输入是什么？", "规则是什么？", "输出是什么？"]
    for i, item in enumerate(left):
        add_card(slide, item, ["药物处理后的样本", "未处理或基线样本", "先由统计规则筛出"][i], 1.0, 2.0 + i * 1.05, 4.2, 0.78, [P.secondary, P.accent, P.green][i])
    for i, item in enumerate(right):
        add_card(slide, item, spec.bullets[i], 6.3, 2.0 + i * 1.05, 4.8, 0.78, [P.secondary, P.accent, P.warning][i])
    add_note(slide, "学生先完成三问，再允许进入工具或 AI 协作。")


def render_matrix_metadata(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    matrix = [
        ["gene", "S1", "S2", "S3", "S4"],
        ["GENE_A", "100", "112", "210", "230"],
        ["GENE_B", "80", "78", "83", "85"],
        ["GENE_C", "240", "228", "90", "95"],
    ]
    metadata = [
        ["sample_id", "group"],
        ["S1", "control"],
        ["S2", "control"],
        ["S3", "treated"],
        ["S4", "treated"],
    ]
    add_simple_table(slide, matrix, 0.95, 2.0, 5.7, 2.3)
    add_simple_table(slide, metadata, 7.1, 2.0, 3.6, 2.3, header_fill=P.secondary)
    add_line(slide, 6.75, 3.15, 7.0, 3.15, color=P.green, width=4)
    add_textbox(slide, "检查点：sample_id 一一对应，分组信息无缺失。", 1.0, 5.15, 8.8, 0.28, size=17, bold=True, color=P.accent)
    add_note(slide, spec.body)


def render_process(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    steps = [
        ("counts", "原始计数矩阵"),
        ("normalization", "样本测序深度校正"),
        ("dispersion", "估计变异程度"),
        ("model test", "组间差异检验"),
        ("results", "输出结果表"),
    ]
    for i, (head, body) in enumerate(steps):
        x = 0.95 + i * 2.22
        add_rect(slide, x, 2.6, 1.82, 1.15, P.surface, line=P.border)
        add_label(slide, str(i + 1), x, 2.6, 0.42, [P.secondary, P.accent, P.green, P.warning, P.slate][i])
        add_textbox(slide, head, x + 0.52, 2.75, 1.1, 0.22, size=12, color=P.text, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, body, x + 0.18, 3.18, 1.45, 0.25, size=10, color=P.muted, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            add_line(slide, x + 1.86, 3.15, x + 2.12, 3.15, color=P.border, width=3)
    add_textbox(slide, "课堂定位：理解每一步在流程中的位置，不推导模型公式。", 1.0, 5.05, 8.8, 0.34, size=17, color=P.accent, bold=True)
    add_note(slide, spec.risk_note)


def render_field_table(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    rows = [
        ["字段", "回答的问题", "课堂提醒"],
        ["gene", "是哪一个基因", "基因功能需核验"],
        ["baseMean", "平均表达水平", "低表达要谨慎"],
        ["log2FoldChange", "变化方向和幅度", "效应量不是机制"],
        ["pvalue", "原始显著性", "大量检验时不够"],
        ["padj", "校正后显著性", "仍需结合效应量"],
    ]
    add_simple_table(slide, rows, 1.0, 1.85, 10.6, 3.7)
    add_note(slide, "学生任务：每个字段写一句“能回答什么，不能回答什么”。")


def render_screen_table(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    rows = [
        ["gene", "baseMean", "log2FC", "padj", "课堂标签"],
        ["GENE_A", "180", "1.8", "0.02", "候选上调"],
        ["GENE_B", "510", "0.1", "0.40", "不确定"],
        ["GENE_C", "12", "-2.1", "0.08", "不确定"],
        ["GENE_D", "90", "-1.4", "0.03", "候选下调"],
    ]
    add_simple_table(slide, rows, 1.0, 1.85, 10.6, 3.2)
    add_textbox(slide, "筛选规则：padj < 0.05 且 abs(log2FoldChange) > 1", 1.0, 5.35, 9.2, 0.32, size=18, color=P.accent, bold=True)
    add_note(slide, spec.body)


def render_compare_padj(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    add_card(slide, "只看 pvalue", "更容易留下大量同时检验中的假阳性；不能独立支持候选基因结论。", 1.0, 2.0, 4.7, 2.5, P.warning)
    add_card(slide, "加入 padj", "把多重检验风险纳入判断；仍然要结合效应量、表达水平和后续证据。", 6.5, 2.0, 4.7, 2.5, P.accent)
    add_textbox(slide, "一句话：padj 帮我们在大量基因同时检验时更谨慎地筛选候选。", 1.0, 5.25, 9.7, 0.38, size=18, color=P.accent, bold=True)
    add_note(slide, spec.body)


def render_volcano(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    chart_x, chart_y, chart_w, chart_h = 5.85, 1.55, 6.0, 4.6
    add_rect(slide, chart_x, chart_y, chart_w, chart_h, P.white, line=P.border)
    add_line(slide, chart_x + 0.6, chart_y + chart_h - 0.55, chart_x + chart_w - 0.35, chart_y + chart_h - 0.55, P.slate, 1)
    add_line(slide, chart_x + chart_w / 2, chart_y + chart_h - 0.55, chart_x + chart_w / 2, chart_y + 0.38, P.border, 1)
    add_line(slide, chart_x + 0.6, chart_y + chart_h - 0.55, chart_x + 0.6, chart_y + 0.38, P.slate, 1)
    add_textbox(slide, "log2FoldChange", chart_x + 2.2, chart_y + chart_h - 0.28, 1.8, 0.18, size=9, color=P.muted, align=PP_ALIGN.CENTER)
    add_textbox(slide, "-log10(padj)", chart_x + 0.1, chart_y + 0.42, 1.2, 0.18, size=9, color=P.muted)
    points = [(-2.2, 2.7, P.secondary), (-1.4, 3.1, P.secondary), (-0.2, 1.0, P.muted), (0.3, 0.8, P.muted), (1.8, 3.3, P.warning), (2.4, 2.5, P.warning), (0.9, 1.7, P.green), (-0.8, 1.5, P.green)]
    for x_val, y_val, color in points:
        px = chart_x + chart_w / 2 + x_val * 0.78
        py = chart_y + chart_h - 0.55 - y_val * 0.82
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(px), Inches(py), Inches(0.12), Inches(0.12))
        dot.fill.solid()
        dot.fill.fore_color.rgb = rgb(color)
        dot.line.fill.background()
    add_bullets(slide, spec.bullets, 0.98, 2.0, 4.0, 2.8, size=17)
    add_note(slide, spec.body)


def render_label_cards(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    rows = [
        ("GENE_A", "padj 0.02, log2FC 1.8", "候选上调", P.warning),
        ("GENE_B", "padj 0.40, log2FC 0.1", "不确定", P.muted),
        ("GENE_C", "padj 0.08, log2FC -2.1", "不确定", P.muted),
        ("GENE_D", "padj 0.03, log2FC -1.4", "候选下调", P.secondary),
    ]
    for i, (gene, rule, tag, color) in enumerate(rows):
        x = 1.0 + (i % 2) * 5.55
        y = 2.0 + (i // 2) * 1.65
        add_card(slide, gene, f"{rule}\n标签：{tag}", x, y, 4.8, 1.2, color)
    add_textbox(slide, "学生先给标签，再解释为什么只是候选。", 1.0, 5.55, 8.0, 0.28, size=17, color=P.accent, bold=True)
    add_note(slide, spec.body)


def render_heatmap(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    x0, y0 = 6.0, 1.65
    genes = ["GENE_A", "GENE_B", "GENE_C", "GENE_D"]
    samples = ["C1", "C2", "T1", "T2"]
    matrix = [
        [0.2, 0.1, 0.9, 0.85],
        [0.45, 0.48, 0.50, 0.52],
        [0.82, 0.76, 0.25, 0.2],
        [0.7, 0.72, 0.28, 0.31],
    ]
    add_textbox(slide, "样本", x0 + 1.15, y0 - 0.28, 2.6, 0.18, size=9, color=P.muted, align=PP_ALIGN.CENTER)
    for c, sample in enumerate(samples):
        add_textbox(slide, sample, x0 + 1.15 + c * 0.68, y0, 0.45, 0.18, size=9, color=P.text, align=PP_ALIGN.CENTER)
    for r, gene in enumerate(genes):
        add_textbox(slide, gene, x0, y0 + 0.4 + r * 0.55, 0.85, 0.18, size=9, color=P.text)
        for c, value in enumerate(matrix[r]):
            color = interpolate_heat(value)
            add_rect(slide, x0 + 1.15 + c * 0.68, y0 + 0.35 + r * 0.55, 0.5, 0.42, color, line=P.white)
    add_bullets(slide, spec.bullets, 0.98, 2.0, 4.2, 2.7, size=17)
    add_textbox(slide, "热图看模式，不替代 DESeq2 统计检验。", 5.95, 4.8, 4.9, 0.28, size=16, color=P.accent, bold=True)
    add_note(slide, spec.body)


def interpolate_heat(value: float) -> str:
    # Blue-white-orange, compact deterministic palette for teaching heatmap.
    if value < 0.34:
        return "8EC3D5"
    if value < 0.67:
        return "F3F8FA"
    return "DE8F05"


def render_evidence_flow(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    steps = [
        ("候选基因", "来自筛选表"),
        ("数据库查询", "GO / Reactome / STRING"),
        ("文献核验", "PubMed / 原始论文"),
        ("保守报告", "候选解释 + 待核验"),
    ]
    for i, (head, body) in enumerate(steps):
        x = 0.95 + i * 2.75
        add_card(slide, head, body, x, 2.4, 2.2, 1.25, [P.secondary, P.accent, P.green, P.warning][i])
        if i < len(steps) - 1:
            add_line(slide, x + 2.23, 3.0, x + 2.68, 3.0, P.border, 3)
    add_textbox(slide, "允许写：这些结果提示某类功能线索，仍需数据库或文献核验。", 1.0, 5.25, 9.3, 0.32, size=17, color=P.accent, bold=True)
    add_note(slide, spec.body)


def render_database_compare(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    items = [
        ("GO", "能说：候选功能术语\n不能说：机制已被证明"),
        ("Reactome", "能说：候选通路线索\n不能说：通路一定被激活"),
        ("STRING", "能说：候选互作网络\n不能说：真实互作已验证"),
    ]
    for i, (head, body) in enumerate(items):
        add_card(slide, head, body, 1.0 + i * 3.65, 2.0, 3.05, 2.35, [P.secondary, P.accent, P.green][i])
    add_textbox(slide, "统一写法：需数据库或文献核验。", 1.0, 5.2, 6.8, 0.3, size=18, color=P.warning, bold=True)
    add_note(slide, spec.body)


def render_ai_prompt(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    add_rect(slide, 0.95, 2.0, 5.3, 2.75, P.surface, line=P.border)
    add_textbox(slide, "推荐 Prompt", 1.2, 2.25, 2.0, 0.22, size=15, color=P.accent, bold=True)
    prompt = "请整理差异表达筛选规则和火山图解释模板；不要编造基因功能；所有功能解释都标注“需数据库或文献核验”。"
    add_textbox(slide, prompt, 1.2, 2.75, 4.4, 1.2, size=16, color=P.text)
    add_rect(slide, 6.95, 2.0, 4.6, 2.75, P.surface, line=P.border)
    add_textbox(slide, "AI 输出审计", 7.2, 2.25, 2.0, 0.22, size=15, color=P.accent, bold=True)
    add_bullets(slide, ["是否新增事实", "是否写成确定结论", "是否标注需核验", "是否保留修改记录"], 7.2, 2.72, 3.6, 1.6, size=14)
    add_note(slide, spec.body)


def render_mistake_table(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    rows = [
        ["常见错误", "纠偏表达"],
        ["只看 pvalue", "同时报告 padj、效应量和表达水平"],
        ["log2FC 大就是医学意义大", "效应量需要生物学和实验背景解释"],
        ["火山图证明机制", "火山图只支持候选筛选和图形观察"],
        ["AI 编造基因功能", "未给来源必须标注需核验"],
    ]
    add_simple_table(slide, rows, 1.0, 1.8, 10.6, 3.7, header_fill=P.warning)
    add_note(slide, "学生任务：选一条错误，改写成可接受表达。")


def render_submission(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    for i, item in enumerate(spec.bullets):
        add_card(slide, f"{i + 1}. {item}", ["阈值与标签", "横轴、纵轴、颜色", "候选而非事实", "需核验原因"][i], 1.0 + i * 2.8, 2.2, 2.25, 1.45, [P.secondary, P.accent, P.green, P.warning][i])
    add_textbox(slide, "评价口径：没有核验记录，最终答案视为不完整。", 1.0, 5.25, 8.8, 0.32, size=18, color=P.accent, bold=True)
    add_note(slide, spec.student_action)


def render_claim_gate(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    rows = [
        ["Claim", "Evidence status", "Action"],
        ["GENE_A 上调", "小表支持", "可写候选上调"],
        ["参与某通路", "需数据库核验", "标注候选解释"],
        ["说明药物有效", "无证据", "不得写入结论"],
    ]
    add_simple_table(slide, rows, 5.65, 1.8, 6.2, 3.0, header_fill=P.accent)
    add_card(slide, "结论写入前先过门", "每一句结论都要回答：证据来自哪里？能否支持这个强度的说法？", 0.95, 2.15, 4.4, 2.0, P.warning)
    add_note(slide, spec.body)


def render_exit_card(slide, spec: SlideSpec) -> None:
    set_common_placeholders(slide, spec)
    add_rect(slide, 1.4, 2.25, 9.6, 2.35, P.surface, line=P.border)
    add_textbox(slide, "出口卡", 1.78, 2.62, 1.4, 0.28, size=20, color=P.accent, bold=True)
    add_textbox(slide, "哪一句结论必须标注“需核验”？为什么？", 1.78, 3.18, 7.8, 0.38, size=24, color=P.text, bold=True)
    add_textbox(slide, "下周进入单细胞转录组可视化：继续训练图形观察与解释边界。", 1.78, 4.05, 7.8, 0.25, size=15, color=P.muted)
    add_note(slide, spec.risk_note)


RENDERERS: dict[str, Callable] = {
    "cover": render_cover,
    "chain": render_chain,
    "three_questions": render_three_questions,
    "matrix_metadata": render_matrix_metadata,
    "process": render_process,
    "field_table": render_field_table,
    "screen_table": render_screen_table,
    "compare_padj": render_compare_padj,
    "volcano": render_volcano,
    "label_cards": render_label_cards,
    "heatmap": render_heatmap,
    "evidence_flow": render_evidence_flow,
    "database_compare": render_database_compare,
    "ai_prompt": render_ai_prompt,
    "mistake_table": render_mistake_table,
    "submission": render_submission,
    "claim_gate": render_claim_gate,
    "exit_card": render_exit_card,
}


def build_working_pptx() -> Presentation:
    prs = Presentation(str(SOURCE_TEMPLATE))
    remove_all_slides(prs)
    for spec in SLIDES:
        layout = prs.slide_layouts[spec.layout_idx]
        slide = prs.slides.add_slide(layout)
        set_common_placeholders(slide, spec)
    prs.save(OUT_DIR / "working.pptx")
    return prs


def build_final_pptx() -> None:
    prs = Presentation(str(OUT_DIR / "working.pptx"))
    for spec, slide in zip(SLIDES, prs.slides):
        renderer = RENDERERS[spec.render]
        renderer(slide, spec)
        if spec.no not in {1, 18}:
            add_source_risk(slide, spec)
    prs.save(OUT_DIR / "final.pptx")


def write_outline() -> None:
    lines = [
        "# Week 15 SYSU Official Blue PPT Outline",
        "",
        "| New Slide | Action Title | Content Role | Exhibit/Image | Template File | Template Slide/Layout |",
        "|---:|---|---|---|---|---|",
    ]
    for spec in SLIDES:
        lines.append(
            f"| {spec.no} | {spec.action_title} | {spec.role} | {spec.exhibit} | {SOURCE_TEMPLATE} | layout {spec.layout_idx} / {spec.layout_name} |"
        )
    lines.append("")
    lines.append("Source storyboard: `course/weeks/week_15/ppt_storyboard.md`.")
    (OUT_DIR / "outline.md").write_text("\n".join(lines), encoding="utf-8")


def write_mapping_and_replacements() -> None:
    mapping = {
        "style_id": STYLE_ID,
        "source_template": str(SOURCE_TEMPLATE),
        "style_spec": str(STYLE_SPEC),
        "asset_manifest": str(ASSET_MANIFEST),
        "storyboard": str(STORYBOARD_PATH),
        "evidence_review": str(EVAL_PATH),
        "slides": [
            {
                "new_slide": spec.no,
                "action_title": spec.action_title,
                "content_role": spec.role,
                "exhibit": spec.exhibit,
                "template_file": str(SOURCE_TEMPLATE),
                "template_slide_layout": {"layout_index": spec.layout_idx, "layout_name": spec.layout_name},
            }
            for spec in SLIDES
        ],
    }
    replacements = {
        "slides": [
            {
                "slide": spec.no,
                "title": spec.action_title,
                "body": spec.body,
                "bullets": spec.bullets,
                "teacher_note": spec.teacher_note,
                "student_action": spec.student_action,
                "risk_note": spec.risk_note,
                "render": spec.render,
            }
            for spec in SLIDES
        ]
    }
    (OUT_DIR / "template-mapping.json").write_text(json.dumps(mapping, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT_DIR / "replacements.json").write_text(json.dumps(replacements, ensure_ascii=False, indent=2), encoding="utf-8")


def copy_style() -> None:
    shutil.copy2(STYLE_SPEC, OUT_DIR / "style.json")


def write_qa_notes() -> None:
    prs = Presentation(str(OUT_DIR / "final.pptx"))
    empty_titles: list[int] = []
    for i, slide in enumerate(prs.slides, 1):
        texts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        if not texts:
            empty_titles.append(i)
    notes = [
        "# Week 15 SYSU Official Blue PPT QA Notes",
        "",
        f"- Style ID: `{STYLE_ID}`",
        f"- Source template: `{SOURCE_TEMPLATE}`",
        f"- Storyboard: `{STORYBOARD_PATH}`",
        f"- Evidence review: `{EVAL_PATH}`",
        f"- Slide count: {len(prs.slides)}",
        f"- Slide size: {prs.slide_width} x {prs.slide_height} EMU",
        f"- Working PPTX: `{OUT_DIR / 'working.pptx'}`",
        f"- Final PPTX: `{OUT_DIR / 'final.pptx'}`",
        "- Placeholder coverage: title/date/footer/slide-number placeholders filled where exposed by the official layouts; content placeholders intentionally cleared before custom exhibit insertion.",
        "- Major inserted exhibits: course chain, count matrix/metadata table, DESeq2 flow, result field table, screening table, pvalue/padj comparison, simulated volcano plot, simulated heatmap, database-boundary comparison, AI prompt audit, claim-evidence gate.",
        "- Simulation boundary: all classroom tables and charts are marked as teaching simulations, not real biomedical conclusions.",
        f"- Empty text slide check: {'none' if not empty_titles else empty_titles}",
        "- Rendering QA: pending PowerPoint COM export.",
        "- Known limitations: generated exhibits are teaching schematics; true biological interpretation requires database or literature verification.",
        "",
    ]
    (OUT_DIR / "qa-notes.md").write_text("\n".join(notes), encoding="utf-8")


def build() -> None:
    ensure_inputs()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    copy_style()
    write_outline()
    write_mapping_and_replacements()
    build_working_pptx()
    build_final_pptx()
    write_qa_notes()
    print(f"built\t{(OUT_DIR / 'final.pptx').relative_to(ROOT)}")
    print(f"qa\t{(OUT_DIR / 'qa-notes.md').relative_to(ROOT)}")


if __name__ == "__main__":
    build()
