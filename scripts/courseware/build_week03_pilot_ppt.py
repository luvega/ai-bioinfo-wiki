"""Build the Week 03 pilot PPTX deck.

The generated PPTX is an output artifact and is written under outputs/.
The script itself is tracked so the pilot workflow can be repeated.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "outputs" / "ppt" / "week_03_pilot"
PPTX_PATH = OUT_DIR / "week_03_ai_python_pilot.pptx"
NOTES_PATH = OUT_DIR / "week_03_speaker_notes.md"


@dataclass(frozen=True)
class Palette:
    ink: str = "1F2933"
    muted: str = "52606D"
    paper: str = "F7F9FC"
    white: str = "FFFFFF"
    teal: str = "0F766E"
    mint: str = "CCFBF1"
    coral: str = "E11D48"
    amber: str = "F59E0B"
    slate: str = "334155"
    soft_blue: str = "DBEAFE"
    soft_rose: str = "FFE4E6"
    soft_amber: str = "FEF3C7"
    soft_green: str = "DCFCE7"


P = Palette()
WIDE_W = 13.333
WIDE_H = 7.5


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def add_textbox(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: int = 20,
    color: str = P.ink,
    bold: bool = False,
    font: str = "Microsoft YaHei",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.05,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.clear()
    box.text_frame.margin_left = Inches(margin)
    box.text_frame.margin_right = Inches(margin)
    box.text_frame.margin_top = Inches(margin)
    box.text_frame.margin_bottom = Inches(margin)
    box.text_frame.vertical_anchor = valign
    para = box.text_frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_title(slide, title: str, subtitle: str | None = None, dark: bool = False):
    color = P.white if dark else P.ink
    add_textbox(slide, title, 0.55, 0.35, 9.8, 0.58, size=28, color=color, bold=True, margin=0)
    if subtitle:
        add_textbox(slide, subtitle, 0.58, 0.93, 9.2, 0.35, size=12, color=P.muted if not dark else P.mint, margin=0)


def add_footer(slide, idx: int):
    add_textbox(slide, f"Week 03 | AI 辅助编程与 Python 快速入门 | {idx:02d}", 0.55, 7.08, 5.0, 0.22, size=8, color=P.muted, margin=0)


def set_bg(slide, color: str = P.paper):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


def add_rect(slide, x: float, y: float, w: float, h: float, fill: str, line: str | None = None, radius: bool = False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    return shp


def add_card(slide, x: float, y: float, w: float, h: float, title: str, body: str, fill: str = P.white, accent: str = P.teal):
    add_rect(slide, x, y, w, h, fill, line="E5E7EB")
    add_rect(slide, x, y, 0.08, h, accent)
    add_textbox(slide, title, x + 0.18, y + 0.18, w - 0.3, 0.28, size=15, bold=True, color=P.ink, margin=0)
    add_textbox(slide, body, x + 0.18, y + 0.55, w - 0.35, h - 0.7, size=11, color=P.slate, margin=0)


def add_bullet_list(slide, items: list[str], x: float, y: float, w: float, h: float, size: int = 13, color: str = P.ink):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = item
        para.level = 0
        para.font.name = "Microsoft YaHei"
        para.font.size = Pt(size)
        para.font.color.rgb = rgb(color)
        para.space_after = Pt(4)
    return box


def add_step(slide, number: str, label: str, x: float, y: float, w: float, color: str):
    add_rect(slide, x, y, w, 0.8, P.white, line="D9E2EC")
    add_rect(slide, x, y, 0.62, 0.8, color)
    add_textbox(slide, number, x + 0.18, y + 0.2, 0.28, 0.3, size=15, color=P.white, bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_textbox(slide, label, x + 0.78, y + 0.18, w - 0.9, 0.38, size=15, bold=True, color=P.ink, margin=0)


def add_code(slide, code: str, x: float, y: float, w: float, h: float, size: int = 14):
    add_rect(slide, x, y, w, h, "111827", line="374151")
    add_textbox(slide, code, x + 0.22, y + 0.18, w - 0.42, h - 0.32, size=size, color="E5E7EB", font="Consolas", margin=0)


def add_table(slide, rows: list[list[str]], x: float, y: float, w: float, h: float, col_widths: list[float] | None = None):
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    if col_widths:
        for i, width in enumerate(col_widths):
            table.columns[i].width = Inches(width)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = text
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            fill = P.teal if r == 0 else P.white
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
                for run in p.runs:
                    run.font.name = "Microsoft YaHei"
                    run.font.size = Pt(10 if r else 9)
                    run.font.bold = r == 0
                    run.font.color.rgb = rgb(P.white if r == 0 else P.ink)
    return table


def build_deck() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(WIDE_W)
    prs.slide_height = Inches(WIDE_H)
    blank = prs.slide_layouts[6]

    # 1
    slide = prs.slides.add_slide(blank)
    set_bg(slide, P.ink)
    add_rect(slide, 0, 0, 13.333, 7.5, P.ink)
    add_rect(slide, 0, 5.95, 13.333, 1.55, P.teal)
    add_textbox(slide, "AI 辅助编程与 Python 快速入门", 0.75, 1.05, 10.2, 0.75, size=34, color=P.white, bold=True, margin=0)
    add_textbox(slide, "医药数据处理与可视化 · Week 03 · 2 学时试点样稿", 0.8, 1.9, 7.8, 0.35, size=14, color=P.mint, margin=0)
    for i, label in enumerate(["问题", "小数据", "代码", "AI 审计", "人工核验"]):
        x = 0.85 + i * 2.35
        add_rect(slide, x, 4.4, 1.55, 0.68, P.white)
        add_textbox(slide, label, x, 4.58, 1.55, 0.24, size=13, color=P.ink, bold=True, align=PP_ALIGN.CENTER, margin=0)
        if i < 4:
            add_textbox(slide, "→", x + 1.67, 4.5, 0.35, 0.25, size=18, color=P.mint, bold=True, margin=0)
    add_textbox(slide, "本样稿目标：用真实页面验证从 Markdown 试点稿到 PPT 试讲稿的工作流", 0.85, 6.28, 9.6, 0.32, size=15, color=P.white, bold=True, margin=0)

    # 2
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "从药学问题到判断规则", "代码不是抽象语法；它把判断规则写成可重复步骤")
    add_table(slide, [["样本", "指标", "状态"], ["S01", "5.2", "正常"], ["S02", "None", "缺失"], ["S03", "12.8", "高风险"]], 0.75, 1.55, 4.6, 2.2, [1.3, 1.5, 1.8])
    add_card(slide, 6.0, 1.35, 5.8, 1.05, "自然语言规则", "忽略缺失值；计算有效值均值；标记大于 7.0 的高风险值。", fill=P.soft_amber, accent=P.amber)
    add_card(slide, 6.0, 2.75, 5.8, 1.05, "代码表达", "把规则拆成输入、处理步骤和输出，后续才能让 AI 帮忙解释或调试。", fill=P.soft_green, accent=P.teal)
    add_card(slide, 6.0, 4.15, 5.8, 1.05, "人工核验", "数据足够小时，先手算预期结果，再判断代码和 AI 输出是否可信。", fill=P.soft_rose, accent=P.coral)
    add_footer(slide, 2)

    # 3
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "Python 在本课程中的位置", "先读懂 10 行以内代码，再进入清洗、统计、绘图与自动化")
    labels = ["数据", "规则", "Python", "图表", "结论"]
    colors = [P.soft_blue, P.soft_amber, P.mint, P.soft_green, P.soft_rose]
    for i, label in enumerate(labels):
        add_rect(slide, 0.8 + i * 2.35, 2.45, 1.55, 1.05, colors[i], line="CBD5E1")
        add_textbox(slide, label, 0.8 + i * 2.35, 2.75, 1.55, 0.28, size=18, bold=True, align=PP_ALIGN.CENTER, margin=0)
        if i < len(labels) - 1:
            add_textbox(slide, "→", 2.55 + i * 2.35, 2.78, 0.38, 0.3, size=20, color=P.teal, bold=True, margin=0)
    add_card(slide, 1.0, 4.55, 3.3, 1.05, "AI 可协助", "解释代码、定位报错、生成测试用例。", fill=P.white, accent=P.teal)
    add_card(slide, 5.0, 4.55, 3.3, 1.05, "学生必须做", "判断输入、输出、边界情况。", fill=P.white, accent=P.coral)
    add_card(slide, 9.0, 4.55, 3.3, 1.05, "教师要检查", "结果是否能由小数据复核。", fill=P.white, accent=P.amber)
    add_footer(slide, 3)

    # 4
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "五个最小语法单元", "够用即可：围绕医药数据处理的最小组合")
    syntax = [
        ("变量", "保存一个指标值\n如 glucose = 5.2", P.soft_blue, P.teal),
        ("列表", "保存一组观测值\n如 [5.2, None, 12.8]", P.soft_green, P.teal),
        ("字典", "保存样本到分组\n如 S01 → control", P.soft_amber, P.amber),
        ("条件", "判断缺失和阈值\n如 x is not None", P.soft_rose, P.coral),
        ("循环", "批量处理多个样本\n如 for x in values", P.mint, P.teal),
    ]
    for i, (title, body, fill, accent) in enumerate(syntax):
        x = 0.65 + (i % 3) * 4.1
        y = 1.55 + (i // 3) * 2.15
        add_card(slide, x, y, 3.45, 1.45, title, body, fill=fill, accent=accent)
    add_textbox(slide, "课堂目标不是记住语法百科，而是能把一个小规则读懂、运行并核验。", 0.9, 6.25, 10.5, 0.35, size=16, bold=True, color=P.teal, margin=0)
    add_footer(slide, 4)

    # 5
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "血糖列表与缺失值", "先识别输入，再讨论代码")
    add_code(slide, "glucose = [5.2, 6.1, 4.9, None, 12.8, 5.7]", 0.8, 1.45, 7.6, 0.85, size=16)
    add_card(slide, 0.95, 3.0, 3.2, 1.35, "缺失值", "`None` 不能直接参与求和，必须先过滤。", fill=P.soft_rose, accent=P.coral)
    add_card(slide, 4.75, 3.0, 3.2, 1.35, "高风险值", "`12.8` 不应删除，而应按阈值标记。", fill=P.soft_amber, accent=P.amber)
    add_card(slide, 8.55, 3.0, 3.2, 1.35, "阈值规则", "课堂采用 `> 7.0` 作为演示阈值。", fill=P.soft_green, accent=P.teal)
    add_textbox(slide, "学生任务：先用笔圈出 None 和 12.8，再写自然语言规则。", 1.0, 5.5, 8.8, 0.32, size=17, bold=True, color=P.ink, margin=0)
    add_footer(slide, 5)

    # 6
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "代码阅读三问", "读代码前先问目标，不从语法细节陷进去")
    add_card(slide, 0.8, 1.55, 3.65, 3.8, "输入是什么", "`glucose` 列表\n包含数值、缺失值和一个高风险值。", fill=P.white, accent=P.teal)
    add_card(slide, 4.85, 1.55, 3.65, 3.8, "规则是什么", "过滤 `None`\n计算有效值均值\n筛选 `> 7.0` 的数值。", fill=P.white, accent=P.amber)
    add_card(slide, 8.9, 1.55, 3.65, 3.8, "输出是什么", "有效数值列表\n均值 `6.94`\n高风险值 `[12.8]`。", fill=P.white, accent=P.coral)
    add_footer(slide, 6)

    # 7
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "最小 Python 实现", "每一行都要能对应一个规则")
    code = "\n".join([
        "valid = [x for x in glucose if x is not None]",
        "mean_glucose = sum(valid) / len(valid)",
        "high_risk = [x for x in valid if x > 7.0]",
    ])
    add_code(slide, code, 0.75, 1.35, 7.7, 2.2, size=15)
    add_step(slide, "1", "过滤缺失", 9.0, 1.35, 3.2, P.teal)
    add_step(slide, "2", "计算均值", 9.0, 2.45, 3.2, P.amber)
    add_step(slide, "3", "筛高风险", 9.0, 3.55, 3.2, P.coral)
    add_card(slide, 0.95, 4.35, 11.2, 0.95, "讲授重点", "代码不是黑箱：列表推导式、sum、len、阈值判断都要回到输入和输出。", fill=P.soft_blue, accent=P.teal)
    add_footer(slide, 7)

    # 8
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "人工核验结果", "先手算，后运行；先预期，后相信")
    add_table(slide, [["项目", "预期结果", "核验方式"], ["有效数值", "[5.2, 6.1, 4.9, 12.8, 5.7]", "排除 None"], ["有效样本数", "5", "手数长度"], ["均值", "6.94", "34.7 / 5"], ["高风险值", "[12.8]", "> 7.0"]], 0.75, 1.35, 7.0, 3.05, [1.55, 3.05, 2.4])
    add_rect(slide, 8.35, 1.45, 3.8, 2.95, P.ink)
    add_textbox(slide, "6.94", 8.35, 1.85, 3.8, 0.85, size=44, color=P.white, bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_textbox(slide, "有效值均值", 8.35, 2.85, 3.8, 0.3, size=16, color=P.mint, bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_textbox(slide, "[12.8]", 8.35, 3.4, 3.8, 0.38, size=21, color=P.white, bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_textbox(slide, "高风险值", 8.35, 3.82, 3.8, 0.25, size=12, color=P.mint, align=PP_ALIGN.CENTER, margin=0)
    add_footer(slide, 8)

    # 9
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "AI 解释代码：先解释，不代写", "合格提示词必须要求 AI 检查缺失值处理")
    prompt = "请逐行解释下面这段 Python 代码的作用，并指出它是否正确处理了 None 缺失值。请先解释，不要直接重写完整代码。"
    add_card(slide, 0.75, 1.45, 6.1, 2.0, "推荐提示词", prompt, fill=P.white, accent=P.teal)
    add_card(slide, 7.25, 1.45, 4.95, 2.0, "审计重点", "是否提到 `None`\n是否说明输入与输出\n是否避免直接改写代码", fill=P.soft_green, accent=P.teal)
    add_bullet_list(slide, ["AI 输出只能作为候选解释", "学生必须用小数据核验", "医学含义不能由 AI 直接判断"], 1.0, 4.35, 10.8, 1.25, size=16, color=P.ink)
    add_footer(slide, 9)

    # 10
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "AI 调试与测试", "错误信息和边界样本同样是课程内容")
    add_code(slide, "mean_glucose = sum(glucose) / len(glucose)\n# TypeError: unsupported operand type(s) for +", 0.75, 1.35, 6.25, 1.45, size=14)
    add_table(slide, [["测试样本", "预期输出"], ["[5.0, 6.0]", "均值 5.5；无高风险"], ["[None, 8.0]", "均值 8.0；高风险 [8.0]"], ["[None]", "无有效值；需提示无法计算"]], 7.45, 1.35, 4.9, 2.7, [2.35, 2.55])
    add_card(slide, 0.95, 4.35, 11.0, 0.95, "讲授重点", "不要让 AI 直接重写一大段。先让它解释错误原因，再由学生补边界测试。", fill=P.soft_amber, accent=P.amber)
    add_footer(slide, 10)

    # 11
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "生信扩展：序列长度与 GC 含量", "只识别数据结构，不要求本周安装 Biopython")
    add_code(slide, "seq = 'ATGCGTAC'\nlength = len(seq)\ngc = (seq.count('G') + seq.count('C')) / length", 0.75, 1.35, 6.1, 1.7, size=15)
    add_card(slide, 7.3, 1.35, 4.9, 1.05, "同一个思路", "输入、规则、输出、核验。", fill=P.soft_green, accent=P.teal)
    add_card(slide, 7.3, 2.75, 4.9, 1.05, "专业数据结构", "DNA、RNA、蛋白质序列也可以被代码处理。", fill=P.soft_blue, accent=P.teal)
    add_card(slide, 7.3, 4.15, 4.9, 1.05, "课堂边界", "这里只讲直觉，不做环境安装。", fill=P.soft_rose, accent=P.coral)
    add_footer(slide, 11)

    # 12
    slide = prs.slides.add_slide(blank)
    set_bg(slide, P.ink)
    add_title(slide, "小结：会读、会测、会问", "能核验的人，才适合让 AI 帮忙写代码", dark=True)
    add_step(slide, "1", "会读：输入、输出、规则", 1.0, 2.1, 3.6, P.teal)
    add_step(slide, "2", "会测：小数据、边界值、预期结果", 4.88, 2.1, 3.95, P.amber)
    add_step(slide, "3", "会问：解释、调试、测试", 9.12, 2.1, 3.45, P.coral)
    add_rect(slide, 1.1, 4.25, 11.0, 1.0, P.white)
    add_textbox(slide, "出口卡：写下一个药学指标规则，并说明你会如何核验 AI 生成的代码。", 1.35, 4.55, 10.4, 0.32, size=18, color=P.ink, bold=True, margin=0)
    add_textbox(slide, "下一步：把这套 PPT 样稿用于真实页面密度与课堂节奏验证。", 1.1, 6.55, 8.5, 0.28, size=12, color=P.mint, margin=0)

    prs.save(PPTX_PATH)
    NOTES_PATH.write_text(
        """# Week 03 PPT 样稿讲稿备注

生成文件：`week_03_ai_python_pilot.pptx`

## 课堂节奏

- 0-10 分钟：问题导入，强调代码是规则的可重复表达。
- 10-25 分钟：五个最小语法单元。
- 25-60 分钟：血糖列表、缺失过滤、均值、高风险标记。
- 60-90 分钟：AI 解释代码和错误调试。
- 90-105 分钟：边界测试样本。
- 105-120 分钟：出口卡与小结。

## 关键核验

- 有效数值：`[5.2, 6.1, 4.9, 12.8, 5.7]`
- 均值：`6.94`
- 高风险值：`[12.8]`
- 典型错误：直接 `sum(glucose)` 会因为 `None` 报错。
""",
        encoding="utf-8",
    )
    print(PPTX_PATH)
    print(NOTES_PATH)


if __name__ == "__main__":
    build_deck()
